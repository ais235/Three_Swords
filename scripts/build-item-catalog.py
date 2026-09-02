from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "src" / "game-data.js"
OUT = ROOT / "artifacts" / "Каталог_предметов_Three_Blades.docx"
ITEM_BASE = ROOT / "assets" / "items"

INK = "171522"
NAVY = "171A33"
GOLD = "D7A928"
PALE_GOLD = "F5E8B7"
PARCHMENT = "FBF8EF"
MUTED = "686475"
BLUE = "2E74B5"
RARE = "2687C8"
EPIC = "8E55C7"
COMMON = "667080"

RU_NAMES = {
    'Small HP Potion +5':'Малое зелье здоровья +5', 'HP Potion +10':'Зелье здоровья +10',
    'HP Potion +15':'Большое зелье здоровья +15', 'Mana Elixir +5':'Эликсир маны +5',
    'Magic Stone (+1 mage attack)':'Магический камень', 'Steel Sword':'Стальной меч',
    'Poison Blade':'Клинок яда', 'Thunder Hammer':'Громовой молот', 'Dark Blade':'Клинок тьмы',
    'Hunter Bow':'Охотничий лук', 'Shadow Bow':'Лук теней', 'Moon Bow':'Лунный лук',
    'War Crossbow':'Боевой арбалет', 'Magic Wand':'Магический жезл', 'Fire Staff':'Огненный посох',
    'Chaos Staff':'Посох хаоса', 'Ice Demon Staff':'Посох ледяного демона',
    'Leather Armor':'Кожаная броня', 'Shadow Cloak':'Плащ теней', 'Runic Shield':'Рунический щит',
    'Mirror Robe':'Мантия отражения', 'Guardian Plate':'Доспех стражника', 'Phoenix Robe':'Одеяние феникса',
    'Dark Mantle':'Тёмная мантия', 'Lucky Amulet':'Амулет удачи', 'Mana Crystal':'Кристалл маны',
    'Scout Eye':'Глаз провидца', 'Hourglass':'Песочные часы', 'Death Ring':'Кольцо смерти',
    'Portal Stone':'Портальный камень', 'Crown of Power':'Корона власти', 'Ash Sword':'Пепельный меч',
    'Forge Hammer':'Молот кузни', 'Coal Bow':'Угольный лук', 'Siege Crossbow':'Осадный арбалет',
    'Magma Wand':'Магмовый жезл', 'Volcano Staff':'Посох вулкана', 'Tide Blade':'Клинок прилива',
    'Abyss Staff':'Посох бездны', 'Vent Shield':'Щит жерла', 'Spark Cloak':'Искровой плащ',
    'Forge Hourglass':'Часы кузни', 'Guardian Trident':'Трезубец хранителя', 'Deep Shell':'Глубинный панцирь',
    'Current Bow':'Лук течений', 'Harpoon Crossbow':'Гарпунный арбалет', 'Sea Ghost Cloak':'Плащ морского призрака',
    'Tide Staff':'Посох прилива', 'Deep Robe':'Глубинная мантия', 'Tide Heart':'Сердце прилива'
}

DESC_RU = {
    'potion_small':'Восстанавливает выбранному герою 5 HP.', 'potion_mid':'Восстанавливает выбранному герою 10 HP.',
    'potion':'Восстанавливает выбранному герою 15 HP.', 'elixir':'Восстанавливает магу 5 MP.',
    'mag_stone':'Навсегда увеличивает базовую атаку Сэрис на 1.',
    'poison_blade':'При попадании отравляет цель на 2 хода.', 'thunder_hammer':'При попадании оглушает врага на 1 ход.',
    'shadow_bow':'Стрелы игнорируют укрытия, лес и союзников на линии огня.',
    'moon_bow':'Первый выстрел каждого раунда наносит в 1,5 раза больше урона.',
    'fire_staff':'Увеличивает радиус всех заклинаний на 1 клетку.',
    'chaos_staff':'После заклинания создаёт случайный взрыв на 5 урона рядом с ближайшими целями.',
    'staff_ice':'Заморозка длится 2 хода вместо 1.', 'shadow_cloak':'Даёт 20% шанс полностью уклониться от атаки.',
    'runic_shield':'С вероятностью 25% отражает негативный статус обратно в атакующего.',
    'mirror_robe':'Возвращает атакующему 10% полученного физического урона.',
    'guardian_plate':'Каждая атака по владельцу наносит атакующему 1 ответный урон.',
    'phoenix':'Один раз за бой воскрешает владельца с 5 HP.', 'phoenix_robe_2':'Один раз за бой воскрешает владельца с 5 HP.',
    'luck_amulet':'Увеличивает золото за каждый бой на 15%.', 'mana_crystal':'В начале каждого хода восстанавливает магу 3 MP.',
    'seer_eye':'На карте раскрывает точный состав врагов выбранной миссии.',
    'hourglass':'Один раз за бой после действия владельца даёт ему дополнительный ход.',
    'death_ring':'При гибели героя создаёт взрыв 3×3 и наносит врагам 15 урона.',
    'portal_stone':'Один раз за бой телепортирует владельца на любую свободную клетку.',
    'crown':'Все остальные герои получают +2 к атаке и +1 к броне.',
    'forge_hammer':'Мощное оружие, оглушающее врага при попадании.', 'volcano_staff':'Увеличивает радиус заклинаний на 1 клетку.',
    'vent_shield':'Атакующий получает 1 ответный урон.', 'spark_cloak':'Даёт 20% шанс уклониться от атаки.',
    'sea_cloak':'Даёт 25% шанс уклониться от атаки.', 'tide_staff':'Заморозка длится 2 хода вместо 1.'
}

MOD_LABELS = {
    'attack':'Атака', 'armor':'Броня', 'range':'Дальность', 'move':'Перемещение', 'maxHp':'Макс. HP',
    'mana':'Мана', 'manaRegen':'Регенерация маны/ход', 'goldBonus':'Золото за бой', 'dodge':'Уклонение',
    'poisonOnHit':'Яд при ударе', 'stunOnHit':'Оглушение при ударе', 'ignoreCover':'Игнор укрытий',
    'firstShotBonus':'Бонус первого выстрела', 'aoePlus':'Радиус заклинаний', 'chaosAoe':'Взрыв после заклинания',
    'freeze2':'Длительность заморозки', 'reflectStatus':'Отражение статуса', 'reflectDmg':'Отражение урона',
    'reflectFlat':'Ответный урон', 'revive5':'Воскрешение', 'extraAction':'Дополнительный ход',
    'deathExplode':'Взрыв при гибели', 'battleTeleport':'Телепортация', 'previewEnemies':'Разведка врагов', 'allyAttack':'Атака союзников',
    'allyArmor':'Броня союзников'
}

def parse_value(raw):
    raw = raw.strip()
    if raw in ('true','false'): return raw == 'true'
    try: return float(raw) if '.' in raw else int(raw)
    except ValueError: return raw.strip("'\"")

def parse_items():
    text = SOURCE.read_text(encoding='utf-8')
    pool = text.split('const ITEM_POOL = [',1)[1].split('];',1)[0]
    items = []
    for obj in re.findall(r"\{\s*id\s*:\s*'[^']+'.*?\}\s*,?", pool):
        def field(name, default=''):
            m = re.search(rf"\b{name}\s*:\s*'([^']*)'", obj)
            return m.group(1) if m else default
        def number(name, default=0):
            m = re.search(rf"\b{name}\s*:\s*(\d+)", obj)
            return int(m.group(1)) if m else default
        mods_m = re.search(r"mods\s*:\s*\{([^}]*)\}", obj)
        mods = {}
        if mods_m:
            for part in mods_m.group(1).split(','):
                if ':' in part:
                    k,v = part.split(':',1); mods[k.strip()] = parse_value(v)
        items.append({'id':field('id'), 'slot':field('slot'), 'heroClass':field('heroClass'),
                      'rarity':field('rarity'), 'minRegion':number('minRegion',1), 'name':field('name'),
                      'price':number('price'), 'sellPrice':number('sellPrice'), 'mods':mods, 'desc':field('desc')})
    shop = text.split('const SHOP_ITEMS = [',1)[1].split('];',1)[0]
    consumables = []
    for line in shop.splitlines():
        if "id:" not in line: continue
        fid = re.search(r"id:\s*'([^']+)'", line).group(1)
        name = re.search(r"name:\s*'([^']+)'", line).group(1)
        price = int(re.search(r"price:\s*(\d+)", line).group(1))
        consumables.append({'id':fid,'slot':'consumable','heroClass':'all','rarity':'common','minRegion':1,
                            'name':name,'price':price,'sellPrice':0,'mods':{},'desc':''})
    icon_block = text.split('const ITEM_ICON_MAP = {',1)[1].split('};',1)[0]
    icons = dict(re.findall(r"\b(\w+)\s*:\s*'([^']+\.png)'", icon_block))
    for item in consumables + items: item['icon'] = icons.get(item['id'],'')
    return consumables + items

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tcPr = cell._tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = mar.find(qn('w:'+tag))
        if node is None: node = OxmlElement('w:'+tag); mar.append(node)
        node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); el = OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def set_font(run, name='Calibri', size=11, color=INK, bold=False, italic=False):
    run.font.name=name; run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),name); run._element.rPr.rFonts.set(qn('w:hAnsi'),name)
    run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string(color); run.bold=bold; run.italic=italic

def add_text(cell, text, size=9.2, color=INK, bold=False, after=2, align=WD_ALIGN_PARAGRAPH.LEFT):
    p=cell.add_paragraph() if cell.text else cell.paragraphs[0]; p.alignment=align
    p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.08
    r=p.add_run(text); set_font(r,size=size,color=color,bold=bold); return p

def mods_text(item):
    out=[]
    for key,val in item['mods'].items():
        label=MOD_LABELS.get(key,key)
        if isinstance(val,bool):
            if key=='aoePlus': shown='+1 клетка'
            elif key=='chaosAoe': shown='5 урона'
            elif key=='freeze2': shown='2 хода'
            elif key=='revive5': shown='5 HP'
            elif key=='extraAction': shown='1 раз за бой'
            elif key=='battleTeleport': shown='1 раз за бой'
            elif key=='previewEnemies': shown='точный состав на карте'
            else: shown='да'
        elif key in ('goldBonus','dodge','reflectStatus','reflectDmg','firstShotBonus'):
            shown=f"{round(val*100)}%"
        elif key=='deathExplode': shown=f"{val} урона"
        else: shown=f"{val:+g}"
        out.append(f"{label}: {shown}")
    return ' · '.join(out) if out else 'Одноразовый эффект'

def role_text(item):
    return {'warrior':'Гарет · воин','archer':'Алира · лучник','mage':'Сэрис · маг','all':'Все герои'}.get(item['heroClass'],'Все герои')

def slot_text(slot): return {'weapon':'Оружие','armor':'Броня','artifact':'Артефакт','consumable':'Расходник'}[slot]
def rarity_text(r): return {'common':'Обычный','rare':'Редкий','epic':'Эпический'}[r]
def rarity_color(r): return {'common':COMMON,'rare':RARE,'epic':EPIC}[r]

def item_desc(item):
    if item['id'] in DESC_RU: return DESC_RU[item['id']]
    return mods_text(item) + '.'

def configure_doc(doc):
    sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=sec.footer_distance=Inches(.492)
    styles=doc.styles
    normal=styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11); normal.font.color.rgb=RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')
    normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.25
    for name,size,before,after,color in [('Heading 1',16,18,10,GOLD),('Heading 2',13,14,7,BLUE),('Heading 3',12,10,5,NAVY)]:
        st=styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
        st._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); st._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')
        st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
    header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.CENTER
    set_font(header.add_run('THREE BLADES  ·  АРСЕНАЛ'),size=8.5,color=MUTED,bold=True)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run('Каталог предметов • актуальные игровые данные'),size=8,color=MUTED)

def add_cover(doc, total):
    for _ in range(5): doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(14)
    set_font(p.add_run('⚔'),name='Segoe UI Symbol',size=26,color=GOLD)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(4)
    set_font(p.add_run('THREE BLADES'),name='Georgia',size=30,color=NAVY,bold=True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(18)
    set_font(p.add_run('АРСЕНАЛ ТРЁХ ЗЕМЕЛЬ'),name='Georgia',size=15,color=GOLD,bold=True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(36)
    set_font(p.add_run('Иллюстрированный каталог оружия, брони, артефактов и расходников'),size=11,color=MUTED,italic=True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(f'{total} предмет  ·  3 героя  ·  3 региона'),size=11,color=NAVY,bold=True)
    doc.add_page_break()

def add_intro(doc, items):
    doc.add_heading('Как читать каталог',level=1)
    p=doc.add_paragraph('Значения перенесены непосредственно из текущих игровых данных. Цена продажи указана только для снаряжения; расходники используются сразу и не продаются.')
    p.paragraph_format.space_after=Pt(12)
    for title,copy in [('Обычный','Базовые предметы, доступные с ранних миссий.'),('Редкий','Сильные предметы с заметным свойством или улучшенными параметрами.'),('Эпический','Ключевые предметы сборки с наиболее сильными эффектами.')]:
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4)
        color={'Обычный':COMMON,'Редкий':RARE,'Эпический':EPIC}[title]
        set_font(p.add_run(title+' — '),size=10.5,color=color,bold=True); set_font(p.add_run(copy),size=10.5,color=INK)
    doc.add_heading('Состав арсенала',level=2)
    counts={s:sum(1 for i in items if i['slot']==s) for s in ('weapon','armor','artifact','consumable')}
    p=doc.add_paragraph(f"Оружие — {counts['weapon']}  ·  Броня — {counts['armor']}  ·  Артефакты — {counts['artifact']}  ·  Расходники — {counts['consumable']}")
    p.paragraph_format.space_after=Pt(12)

def add_catalog_table(doc, title, items):
    doc.add_heading(title,level=1)
    table=doc.add_table(rows=1,cols=3); table.alignment=WD_TABLE_ALIGNMENT.LEFT; table.autofit=False
    widths=[Inches(.95),Inches(2.35),Inches(3.20)]
    table.columns[0].width=widths[0]; table.columns[1].width=widths[1]; table.columns[2].width=widths[2]
    headers=['ОБЛИК','ПРЕДМЕТ','ХАРАКТЕРИСТИКИ И ЭФФЕКТ']
    for idx,cell in enumerate(table.rows[0].cells):
        cell.width=widths[idx]; set_cell_shading(cell,NAVY); set_cell_margins(cell,100,120,100,120)
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.text=''
        add_text(cell,headers[idx],size=8,color='FFFFFF',bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    for item in items:
        cells=table.add_row().cells
        for idx,c in enumerate(cells): c.width=widths[idx]; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(c,110,120,110,120); set_cell_shading(c,PARCHMENT if len(table.rows)%2 else 'FFFFFF')
        cells[0].text=''; p=cells[0].paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        icon=ITEM_BASE/item['icon']
        if icon.exists():
            picture = p.add_run().add_picture(str(icon),width=Inches(.68))
            picture._inline.docPr.set('descr', f"Иконка предмета: {RU_NAMES.get(item['name'], item['name'])}")
            picture._inline.docPr.set('title', RU_NAMES.get(item['name'], item['name']))
        else: set_font(p.add_run('◇'),name='Segoe UI Symbol',size=24,color=GOLD)
        cells[1].text=''; add_text(cells[1],RU_NAMES.get(item['name'],item['name']),size=10.2,color=rarity_color(item['rarity']),bold=True,after=3)
        add_text(cells[1],f"{slot_text(item['slot'])} · {role_text(item)}",size=8.3,color=MUTED,after=2)
        if item['slot']=='consumable': price=f"Покупка: {item['price']} золота"
        else: price=f"Покупка: {item['price']} · Продажа: {item['sellPrice']} золота"
        add_text(cells[1],f"{rarity_text(item['rarity'])} · Земля {item['minRegion']}",size=8.3,color=rarity_color(item['rarity']),bold=True,after=2)
        add_text(cells[1],price,size=8.2,color=MUTED,after=0)
        cells[2].text=''; add_text(cells[2],mods_text(item),size=9,color=NAVY,bold=True,after=4)
        add_text(cells[2],item_desc(item),size=8.6,color=INK,after=0)
    # Exact 9360 DXA geometry with 120 DXA indent.
    tblPr=table._tbl.tblPr; tblW=tblPr.find(qn('w:tblW')); tblW.set(qn('w:w'),'9360'); tblW.set(qn('w:type'),'dxa')
    ind=OxmlElement('w:tblInd'); ind.set(qn('w:w'),'120'); ind.set(qn('w:type'),'dxa'); tblPr.append(ind)
    grid=table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    dxa=[1368,3384,4608]
    for value in dxa:
        col=OxmlElement('w:gridCol'); col.set(qn('w:w'),str(value)); grid.append(col)
    for row in table.rows:
        for cell,value in zip(row.cells,dxa):
            tcW=cell._tc.get_or_add_tcPr().find(qn('w:tcW')); tcW.set(qn('w:w'),str(value)); tcW.set(qn('w:type'),'dxa')
    doc.add_paragraph()

def build():
    items=parse_items(); OUT.parent.mkdir(exist_ok=True)
    doc=Document(); configure_doc(doc); add_cover(doc,len(items)); add_intro(doc,items)
    groups=[
        ('Расходники',[i for i in items if i['slot']=='consumable']),
        ('Оружие Гарета',[i for i in items if i['slot']=='weapon' and i['heroClass']=='warrior']),
        ('Оружие Алиры',[i for i in items if i['slot']=='weapon' and i['heroClass']=='archer']),
        ('Оружие Сэрис',[i for i in items if i['slot']=='weapon' and i['heroClass']=='mage']),
        ('Броня',[i for i in items if i['slot']=='armor']),
        ('Артефакты',[i for i in items if i['slot']=='artifact']),
    ]
    for title,group in groups: add_catalog_table(doc,title,group)
    doc.core_properties.title='Three Blades — Арсенал трёх земель'
    doc.core_properties.subject='Иллюстрированный каталог игровых предметов'
    doc.core_properties.author='Three Blades'
    doc.save(OUT)
    print(f"created={OUT}")
    print(f"items={len(items)} missing_icons={sum(1 for i in items if not (ITEM_BASE/i['icon']).exists())}")

if __name__=='__main__': build()
